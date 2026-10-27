"""Structured text extraction for Markdown, PDF, PowerPoint and Word files."""

from __future__ import annotations

import mimetypes
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import cast

from .chunker import parse_frontmatter
from .vision import VisionExtractor, get_vision_extractor

SUPPORTED_EXTENSIONS = frozenset({".md", ".pdf", ".pptx", ".docx"})
_VISION_UNSET = object()


@dataclass(frozen=True)
class SourceSection:
    heading: str
    text: str
    page_number: int | None = None
    slide_number: int | None = None
    content_type: str = "text"


@dataclass(frozen=True)
class ParsedSource:
    title: str
    category: str
    source: str
    source_file: str
    file_type: str
    version: str
    sections: tuple[SourceSection, ...]


class DocumentParseError(ValueError):
    pass


def _category(path: Path, corpus_root: Path) -> str:
    try:
        return path.relative_to(corpus_root).parts[0]
    except (ValueError, IndexError):
        return "default"


def _relative_source(path: Path, corpus_root: Path) -> str:
    try:
        return path.relative_to(corpus_root).as_posix()
    except ValueError:
        return path.name


def parse_markdown(path: Path, corpus_root: Path) -> ParsedSource:
    meta, body = parse_frontmatter(path.read_text(encoding="utf-8"))
    heading_title = next(
        (line.lstrip("# ").strip() for line in body.splitlines() if line.startswith("# ")),
        None,
    )
    return ParsedSource(
        title=meta.get("title") or heading_title or path.stem,
        category=meta.get("category") or _category(path, corpus_root),
        source=meta.get("source", ""),
        source_file=_relative_source(path, corpus_root),
        file_type="markdown",
        version=meta.get("version", "1"),
        sections=(SourceSection(heading="", text=body),),
    )


def _rows_to_markdown(rows: list[list[str | None]]) -> str:
    cleaned = [[(cell or "").strip().replace("\n", " ") for cell in row] for row in rows if row]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    normalized = [row + [""] * (width - len(row)) for row in cleaned]
    return "\n".join(
        [
            "| " + " | ".join(normalized[0]) + " |",
            "| " + " | ".join(["---"] * width) + " |",
        ]
        + ["| " + " | ".join(row) + " |" for row in normalized[1:]]
    )


def _png_bytes(image) -> bytes:
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def parse_pdf(
    path: Path,
    corpus_root: Path,
    vision_extractor: VisionExtractor | None = None,
) -> ParsedSource:
    import pdfplumber

    sections: list[SourceSection] = []
    title = path.stem
    with pdfplumber.open(path) as document:
        title = (document.metadata or {}).get("Title") or path.stem
        for page_number, page in enumerate(document.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(
                    SourceSection(
                        f"第 {page_number} 页",
                        text,
                        page_number=page_number,
                    )
                )
            for table_number, rows in enumerate(page.extract_tables(), start=1):
                table = _rows_to_markdown(rows)
                if table:
                    sections.append(
                        SourceSection(
                            f"第 {page_number} 页 / 表格 {table_number}",
                            table,
                            page_number=page_number,
                            content_type="table",
                        )
                    )
            if not text and vision_extractor is not None:
                rendered = page.to_image(resolution=vision_extractor.pdf_dpi).original
                extracted = vision_extractor.extract(
                    _png_bytes(rendered),
                    "image/png",
                    context=f"{path.name} 第 {page_number} 页（扫描页 OCR）",
                ).strip()
                if extracted:
                    sections.append(
                        SourceSection(
                            f"第 {page_number} 页 / OCR",
                            extracted,
                            page_number=page_number,
                            content_type="ocr",
                        )
                    )
    return ParsedSource(
        title=title,
        category=_category(path, corpus_root),
        source="",
        source_file=_relative_source(path, corpus_root),
        file_type="pdf",
        version="1",
        sections=tuple(sections),
    )


def _table_markdown(table) -> str:
    return _rows_to_markdown([[cell.text for cell in row.cells] for row in table.rows])


def parse_pptx(
    path: Path,
    corpus_root: Path,
    vision_extractor: VisionExtractor | None = None,
) -> ParsedSource:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    sections: list[SourceSection] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title else f"幻灯片 {slide_number}"
        text_blocks: list[str] = []
        tables: list[str] = []
        images: list[tuple[bytes, str]] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table = _table_markdown(shape.table)
                if table:
                    tables.append(table)
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text and text != title:
                    text_blocks.append(text)
            elif vision_extractor is not None and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images.append((shape.image.blob, shape.image.content_type))
        text = "\n\n".join(text_blocks).strip()
        if text:
            sections.append(SourceSection(title, text, slide_number=slide_number))
        sections.extend(
            SourceSection(
                f"{title} / 表格 {table_number}",
                table,
                slide_number=slide_number,
                content_type="table",
            )
            for table_number, table in enumerate(tables, start=1)
        )
        if vision_extractor is not None:
            for image_number, (image_bytes, media_type) in enumerate(images, start=1):
                extracted = vision_extractor.extract(
                    image_bytes,
                    media_type,
                    context=f"{path.name} 第 {slide_number} 页图片 {image_number}",
                ).strip()
                if extracted:
                    sections.append(
                        SourceSection(
                            f"{title} / 图片 {image_number}",
                            extracted,
                            slide_number=slide_number,
                            content_type="image",
                        )
                    )
    return ParsedSource(
        title=(presentation.core_properties.title or path.stem).strip(),
        category=_category(path, corpus_root),
        source="",
        source_file=_relative_source(path, corpus_root),
        file_type="pptx",
        version="1",
        sections=tuple(sections),
    )


def _docx_images(path: Path) -> list[tuple[str, bytes, str]]:
    images: list[tuple[str, bytes, str]] = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith("word/media/") or name.endswith("/"):
                continue
            media_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
            if media_type.startswith("image/"):
                images.append((Path(name).name, archive.read(name), media_type))
    return images


def parse_docx(
    path: Path,
    corpus_root: Path,
    vision_extractor: VisionExtractor | None = None,
) -> ParsedSource:
    import mammoth

    with path.open("rb") as source:
        result = mammoth.convert_to_markdown(source)
    body = result.value.strip()
    sections: list[SourceSection] = []
    if body:
        sections.append(SourceSection("", body))
    if vision_extractor is not None:
        for image_number, (name, image_bytes, media_type) in enumerate(_docx_images(path), start=1):
            extracted = vision_extractor.extract(
                image_bytes,
                media_type,
                context=f"{path.name} 嵌入图片 {image_number}（{name}）",
            ).strip()
            if extracted:
                sections.append(
                    SourceSection(
                        f"嵌入图片 {image_number}",
                        extracted,
                        content_type="image",
                    )
                )
    return ParsedSource(
        title=path.stem,
        category=_category(path, corpus_root),
        source="",
        source_file=_relative_source(path, corpus_root),
        file_type="docx",
        version="1",
        sections=tuple(sections),
    )


def parse_source(
    path: Path,
    corpus_root: Path,
    *,
    vision_extractor: VisionExtractor | None | object = _VISION_UNSET,
) -> ParsedSource:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(f"Unsupported document type: {path.suffix or '(none)'}")
    try:
        extractor = get_vision_extractor() if vision_extractor is _VISION_UNSET else cast(VisionExtractor | None, vision_extractor)
        if suffix == ".md":
            parsed = parse_markdown(path, corpus_root)
        elif suffix == ".pdf":
            parsed = parse_pdf(path, corpus_root, extractor)
        elif suffix == ".pptx":
            parsed = parse_pptx(path, corpus_root, extractor)
        else:
            parsed = parse_docx(path, corpus_root, extractor)
    except Exception as exc:
        raise DocumentParseError(f"Failed to parse {path.name}: {exc}") from exc
    if not parsed.sections or not any(section.text.strip() for section in parsed.sections):
        raise DocumentParseError(f"No extractable text found in {path.name}")
    return parsed
