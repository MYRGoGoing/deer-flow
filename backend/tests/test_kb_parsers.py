import zipfile
from pathlib import Path

import pytest

from deerflow.community.kb_search.config import ChunkingConfig, KnowledgeBaseConfig
from deerflow.community.kb_search.ingest import chunks_for_source
from deerflow.community.kb_search.parsers import DocumentParseError, parse_source


class FakeVisionExtractor:
    name = "fake-vision"
    pdf_dpi = 72

    def __init__(self, text: str) -> None:
        self.text = text
        self.calls: list[tuple[bytes, str, str]] = []

    def extract(self, image_bytes: bytes, media_type: str, *, context: str = "") -> str:
        self.calls.append((image_bytes, media_type, context))
        return self.text


def _write_png(path: Path) -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (240, 100), "white")
    ImageDraw.Draw(image).text((12, 36), "FAULT CODE E102", fill="black")
    image.save(path, format="PNG")
    return path.read_bytes()


def _write_pdf(path: Path, text: str) -> None:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend((f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n").encode("ascii"))
    path.write_bytes(payload)


def _write_docx(path: Path) -> None:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    relationships = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Access Control</w:t></w:r></w:p>
    <w:p><w:r><w:t>Operators need approval before deployment.</w:t></w:r></w:p>
  </w:body>
</w:document>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def test_parse_markdown_frontmatter_and_relative_source(tmp_path):
    path = tmp_path / "platform" / "guide.md"
    path.parent.mkdir()
    path.write_text(
        "---\ntitle: RBFT Guide\ncategory: consensus\nversion: '2'\nsource: https://docs.example/rbft\n---\n# Overview\nRBFT rotates leaders.",
        encoding="utf-8",
    )

    parsed = parse_source(path, tmp_path)

    assert parsed.title == "RBFT Guide"
    assert parsed.category == "consensus"
    assert parsed.version == "2"
    assert parsed.source_file == "platform/guide.md"
    assert parsed.sections[0].text.startswith("# Overview")


def test_parse_pdf_preserves_page_number_in_chunks(tmp_path):
    path = tmp_path / "platform" / "manual.pdf"
    path.parent.mkdir()
    _write_pdf(path, "Maintenance torque is 120 Nm")

    parsed = parse_source(path, tmp_path)
    chunks = chunks_for_source(
        "pdf-doc",
        parsed,
        _config(tmp_path, min_child_chars=5),
    )

    assert parsed.file_type == "pdf"
    assert parsed.sections[0].page_number == 1
    assert "120 Nm" in parsed.sections[0].text
    assert chunks
    assert {chunk.page_number for chunk in chunks} == {1}
    assert {chunk.source_file for chunk in chunks} == {"platform/manual.pdf"}


def test_parse_scanned_pdf_uses_vision_ocr(tmp_path):
    from PIL import Image

    path = tmp_path / "platform" / "scanned.pdf"
    path.parent.mkdir()
    image = Image.new("RGB", (240, 100), "white")
    image.save(path, format="PDF")
    extractor = FakeVisionExtractor("故障代码 E102：检查主轴温度传感器。")

    parsed = parse_source(path, tmp_path, vision_extractor=extractor)

    assert parsed.sections[0].content_type == "ocr"
    assert parsed.sections[0].page_number == 1
    assert "E102" in parsed.sections[0].text
    assert extractor.calls[0][1] == "image/png"
    assert "扫描页 OCR" in extractor.calls[0][2]
    chunks = chunks_for_source(
        "scanned-pdf",
        parsed,
        _config(tmp_path, min_child_chars=5),
    )
    assert {chunk.content_type for chunk in chunks} == {"ocr"}
    assert {chunk.page_number for chunk in chunks} == {1}


def test_parse_pptx_extracts_text_table_and_slide_number(tmp_path):
    from pptx import Presentation

    path = tmp_path / "baas" / "deployment.pptx"
    path.parent.mkdir()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Deployment Matrix"
    slide.placeholders[1].text = "Use HA mode for production clusters."
    table = slide.shapes.add_table(2, 2, 50, 200, 500, 100).table
    table.cell(0, 0).text = "Mode"
    table.cell(0, 1).text = "Nodes"
    table.cell(1, 0).text = "HA"
    table.cell(1, 1).text = "3"
    presentation.save(path)

    parsed = parse_source(path, tmp_path)

    assert parsed.file_type == "pptx"
    assert parsed.sections[0].heading == "Deployment Matrix"
    assert parsed.sections[0].slide_number == 1
    assert "Use HA mode" in parsed.sections[0].text
    table_section = parsed.sections[1]
    assert table_section.heading == "Deployment Matrix / 表格 1"
    assert table_section.slide_number == 1
    assert table_section.content_type == "table"
    assert "| Mode | Nodes |" in table_section.text


def test_parse_pptx_extracts_embedded_image_with_vision(tmp_path):
    from pptx import Presentation

    image_path = tmp_path / "fault.png"
    _write_png(image_path)
    path = tmp_path / "baas" / "fault-guide.pptx"
    path.parent.mkdir()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 50, 50, width=400, height=160)
    presentation.save(path)
    extractor = FakeVisionExtractor("故障图显示 E102 指向温度传感器。")

    parsed = parse_source(path, tmp_path, vision_extractor=extractor)

    assert len(parsed.sections) == 1
    assert parsed.sections[0].content_type == "image"
    assert parsed.sections[0].slide_number == 1
    assert "E102" in parsed.sections[0].text
    assert extractor.calls[0][1] == "image/png"


def test_parse_docx_extracts_heading_and_body(tmp_path):
    path = tmp_path / "contract-dev" / "policy.docx"
    path.parent.mkdir()
    _write_docx(path)

    parsed = parse_source(path, tmp_path)

    assert parsed.file_type == "docx"
    assert parsed.source_file == "contract-dev/policy.docx"
    assert "# Access Control" in parsed.sections[0].text
    assert "approval before deployment" in parsed.sections[0].text


def test_parse_docx_extracts_embedded_image_with_vision(tmp_path):
    path = tmp_path / "contract-dev" / "illustrated-policy.docx"
    path.parent.mkdir()
    _write_docx(path)
    image_path = tmp_path / "diagram.png"
    image_bytes = _write_png(image_path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("word/media/diagram.png", image_bytes)
    extractor = FakeVisionExtractor("审批流程：申请人提交后由管理员复核。")

    parsed = parse_source(path, tmp_path, vision_extractor=extractor)

    image_sections = [section for section in parsed.sections if section.content_type == "image"]
    assert len(image_sections) == 1
    assert "管理员复核" in image_sections[0].text
    assert "diagram.png" in extractor.calls[0][2]


@pytest.mark.parametrize("filename", ["notes.txt", "README"])
def test_parse_source_rejects_unsupported_type(tmp_path, filename):
    path = tmp_path / filename
    path.write_text("plain text", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="Unsupported document type"):
        parse_source(path, tmp_path)


def test_parse_source_rejects_document_without_extractable_text(tmp_path):
    path = tmp_path / "empty.md"
    path.write_text("---\ntitle: Empty\n---\n", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="No extractable text"):
        parse_source(path, tmp_path)


def _config(root: Path, *, min_child_chars: int) -> KnowledgeBaseConfig:
    from deerflow.community.kb_search.config import (
        EmbeddingConfig,
        EvaluationConfig,
        RerankerConfig,
        RetrievalConfig,
    )

    return KnowledgeBaseConfig(
        project_root=root,
        corpus_dir=root,
        index_dir=root / "index",
        max_document_bytes=1024,
        chunking=ChunkingConfig(min_child_chars=min_child_chars),
        retrieval=RetrievalConfig(),
        embedding=EmbeddingConfig(),
        reranker=RerankerConfig(),
        evaluation=EvaluationConfig(),
    )
